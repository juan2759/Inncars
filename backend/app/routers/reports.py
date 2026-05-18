from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, extract
from typing import Optional
from datetime import date, datetime
import io

from app.database import get_db
from app.models.sale import Sale, SaleStatus
from app.models.vehicle import Vehicle, VehicleStatus
from app.models.customer import Customer
from app.models.workshop import WorkOrder, WorkOrderStatus, WorkOrderItem
from app.auth import get_current_user
from app.models.user import User

router = APIRouter(prefix="/api/reports", tags=["reports"])

@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    today = date.today()
    month_start = today.replace(day=1)

    total_vehicles = db.query(func.count(Vehicle.id)).scalar()
    available = db.query(func.count(Vehicle.id)).filter(Vehicle.estado == VehicleStatus.disponible).scalar()
    reserved = db.query(func.count(Vehicle.id)).filter(Vehicle.estado == VehicleStatus.reservado).scalar()
    sold = db.query(func.count(Vehicle.id)).filter(Vehicle.estado == VehicleStatus.vendido).scalar()

    sales_month = db.query(Sale).filter(
        Sale.estado == SaleStatus.cerrada,
        func.date(Sale.fecha_venta) >= month_start,
    ).all()

    sales_today = db.query(Sale).filter(
        Sale.estado == SaleStatus.cerrada,
        func.date(Sale.fecha_venta) == today,
    ).all()

    wo_pending = db.query(func.count(WorkOrder.id)).filter(
        WorkOrder.estado == WorkOrderStatus.pendiente
    ).scalar()
    wo_in_progress = db.query(func.count(WorkOrder.id)).filter(
        WorkOrder.estado == WorkOrderStatus.en_proceso
    ).scalar()

    workshop_month = db.query(WorkOrderItem).join(
        WorkOrder, WorkOrder.id == WorkOrderItem.work_order_id
    ).filter(
        WorkOrder.estado.in_([WorkOrderStatus.finalizado, WorkOrderStatus.entregado]),
        func.date(WorkOrder.fecha_finalizado) >= month_start,
    ).all()

    new_customers_month = db.query(func.count(Customer.id)).filter(
        func.date(Customer.created_at) >= month_start
    ).scalar()

    return {
        "stock": {
            "total": total_vehicles,
            "disponible": available,
            "reservado": reserved,
            "vendido": sold,
        },
        "ventas": {
            "hoy_cantidad": len(sales_today),
            "hoy_total": sum(s.precio_venta for s in sales_today),
            "mes_cantidad": len(sales_month),
            "mes_total": sum(s.precio_venta for s in sales_month),
        },
        "taller": {
            "pendientes": wo_pending,
            "en_proceso": wo_in_progress,
            "ingresos_mes": sum(i.subtotal for i in workshop_month),
        },
        "clientes": {
            "nuevos_mes": new_customers_month,
            "total": db.query(func.count(Customer.id)).scalar(),
        },
    }

@router.get("/sales-by-period")
def sales_by_period(
    period: str = Query("month", regex="^(day|week|month|year)$"),
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    q = db.query(Sale).filter(Sale.estado == SaleStatus.cerrada)
    if fecha_desde:
        q = q.filter(func.date(Sale.fecha_venta) >= fecha_desde)
    if fecha_hasta:
        q = q.filter(func.date(Sale.fecha_venta) <= fecha_hasta)
    sales = q.all()

    result = {}
    for s in sales:
        if period == "day":
            key = str(s.fecha_venta.date())
        elif period == "week":
            key = f"{s.fecha_venta.isocalendar()[0]}-W{s.fecha_venta.isocalendar()[1]:02d}"
        elif period == "month":
            key = s.fecha_venta.strftime("%Y-%m")
        else:
            key = str(s.fecha_venta.year)
        if key not in result:
            result[key] = {"cantidad": 0, "total": 0}
        result[key]["cantidad"] += 1
        result[key]["total"] += s.precio_venta

    return [{"periodo": k, **v} for k, v in sorted(result.items())]

@router.get("/export/sales")
def export_sales_excel(
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    q = db.query(Sale).filter(Sale.estado == SaleStatus.cerrada)
    if fecha_desde:
        q = q.filter(func.date(Sale.fecha_venta) >= fecha_desde)
    if fecha_hasta:
        q = q.filter(func.date(Sale.fecha_venta) <= fecha_hasta)
    sales = q.order_by(Sale.fecha_venta).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    headers = ["ID", "Fecha", "Vehículo", "Cliente", "Precio", "Forma de Pago", "Estado"]
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row, s in enumerate(sales, 2):
        vehicle = db.query(Vehicle).filter(Vehicle.id == s.vehicle_id).first()
        customer = db.query(Customer).filter(Customer.id == s.customer_id).first()
        ws.cell(row=row, column=1, value=s.id)
        ws.cell(row=row, column=2, value=s.fecha_venta.strftime("%d/%m/%Y") if s.fecha_venta else "")
        ws.cell(row=row, column=3, value=f"{vehicle.marca} {vehicle.modelo} {vehicle.anio}" if vehicle else "")
        ws.cell(row=row, column=4, value=f"{customer.nombre} {customer.apellido}" if customer else "")
        ws.cell(row=row, column=5, value=s.precio_venta)
        ws.cell(row=row, column=6, value=s.forma_pago.value)
        ws.cell(row=row, column=7, value=s.estado.value)

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=ventas.xlsx"},
    )

@router.get("/export/stock")
def export_stock_excel(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    vehicles = db.query(Vehicle).order_by(Vehicle.marca, Vehicle.modelo).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"

    headers = ["ID", "Marca", "Modelo", "Año", "Versión", "Color", "Km", "Patente", "Combustible", "Transmisión", "Precio Costo", "Precio Venta", "Estado"]
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row, v in enumerate(vehicles, 2):
        vals = [v.id, v.marca, v.modelo, v.anio, v.version or "", v.color or "",
                v.kilometraje, v.patente or "", v.tipo_combustible.value,
                v.transmision.value, v.precio_costo, v.precio_venta, v.estado.value]
        for col, val in enumerate(vals, 1):
            ws.cell(row=row, column=col, value=val)

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=stock.xlsx"},
    )
